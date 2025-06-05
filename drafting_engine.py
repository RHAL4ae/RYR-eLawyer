import argparse
from typing import Dict

try:
    from googletrans import Translator
except Exception:
    Translator = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from fpdf import FPDF
except ImportError:
    FPDF = None

LEGAL_TERMS: Dict[str, Dict[str, str]] = {
    'uae': {
        'company': 'شركة',
        'contract': 'عقد',
    },
    'ksa': {
        'company': 'مؤسسة',
        'contract': 'اتفاقية',
    }
}

def translate_text(text: str, dest: str) -> str:
    """Translate text using googletrans if available, otherwise return a placeholder."""
    if Translator is None:
        return f"[Translation unavailable to {dest}] {text}"
    translator = Translator()
    try:
        result = translator.translate(text, dest=dest)
        return result.text
    except Exception as e:
        return f"[Translation Error: {e}] {text}"

def adapt_tone(text: str, tone: str) -> str:
    if tone == 'formal':
        return text.replace('you', 'Sir/Madam')
    elif tone == 'casual':
        return text.replace('Sir/Madam', 'you')
    return text

def map_legal_terms(text: str, jurisdiction: str) -> str:
    mapping = LEGAL_TERMS.get(jurisdiction.lower())
    if not mapping:
        return text
    for eng, local in mapping.items():
        text = text.replace(eng, local)
    return text

def export_word(text: str, output_file: str):
    if Document is None:
        raise RuntimeError('python-docx not installed')
    doc = Document()
    doc.add_paragraph(text)
    doc.save(output_file)

def export_pdf(text: str, output_file: str):
    if FPDF is None:
        raise RuntimeError('fpdf not installed')
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Arial', size=12)
    for line in text.split('\n'):
        pdf.cell(200, 10, txt=line, ln=True)
    pdf.output(output_file)

def main():
    parser = argparse.ArgumentParser(description='Multilingual Legal Drafting Engine')
    parser.add_argument('--text', required=True, help='Source text to draft')
    parser.add_argument('--target_lang', default='ar', help='Language code for translation (e.g., ar, en)')
    parser.add_argument('--tone', choices=['formal', 'casual'], default='formal')
    parser.add_argument('--jurisdiction', default='uae')
    parser.add_argument('--export', choices=['pdf', 'word'], help='Export format')
    parser.add_argument('--output', help='Output file name')
    args = parser.parse_args()

    translated = translate_text(args.text, args.target_lang)
    toned = adapt_tone(translated, args.tone)
    localized = map_legal_terms(toned, args.jurisdiction)

    if args.export == 'word' and args.output:
        export_word(localized, args.output)
    elif args.export == 'pdf' and args.output:
        export_pdf(localized, args.output)
    else:
        print(localized)

if __name__ == '__main__':
    main()
